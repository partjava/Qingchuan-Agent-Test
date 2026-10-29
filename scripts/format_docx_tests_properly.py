# -*- coding: utf-8 -*-
import docx

doc_path = "AI智能体测试挑战赛-晴川校园办事智能体测试与可靠性评估.docx"
doc = docx.Document(doc_path)

# 找到 5.2.1 节
p_521_idx = None
for i, p in enumerate(doc.paragraphs):
    if "5.2.1 代码测试范围与执行入口" in p.text:
        p_521_idx = i
        break

if p_521_idx is not None:
    # 找到 5.2.2 节
    p_522_idx = None
    for j in range(p_521_idx + 1, len(doc.paragraphs)):
        if "5.2.2 用例清单展示" in doc.paragraphs[j].text:
            p_522_idx = j
            break
            
    # 删除之前添加在 5.2.1 和 5.2.2 之间的测试说明段落，重新插入规范排版
    for k in range(p_522_idx - 1, p_521_idx, -1):
        txt = doc.paragraphs[k].text
        if any(marker in txt for marker in ["项目完整开源代码", "tests/ 目录下 25 份", "【开源代码与测试工程仓库】", "一、 智能体与 LangGraph"]):
            p_del = doc.paragraphs[k]._element
            p_del.getparent().remove(p_del)

# 重新定位 5.2.2 节
p_522_idx = None
for j in range(len(doc.paragraphs)):
    if "5.2.2 用例清单展示" in doc.paragraphs[j].text:
        p_522_idx = j
        break

ref_p = doc.paragraphs[p_522_idx]

# 插入开源仓库说明
p_repo = ref_p.insert_paragraph_before()
r_repo = p_repo.add_run("开源代码与测试工程仓库地址：https://github.com/partjava/Qingchuan-Agent-Test")
r_repo.bold = True

# 插入核心主标题
p_main_title = ref_p.insert_paragraph_before()
r_main = p_main_title.add_run("tests/ 目录下 25 份测试脚本说明：")
r_main.bold = True
r_main.font.size = docx.shared.Pt(11)

sections_data = [
    ("一、 智能体与 LangGraph 调度测试（Agent 专项）", 
     "专门测试多智能体决策、工具调用、状态流转和死循环熔断：\n"
     "• test_agent_graph.py：测试基于 LangGraph 构建的状态机图（StateGraph），验证节点流转、条件边分支选择、独立 Token 预算以及 8 步强制递归熔断。\n"
     "• test_agent_planner.py：测试智能体规划器（Planner），模拟真实 HTTP 交互，验证多步骤拆解是否合理。\n"
     "• test_agent_plan.py：测试智能体长链条规划的具体执行逻辑与证据整合。\n"
     "• test_campus_agent.py：智能体行为契约测试，验证智能体在面对校园办事任务时的工具调用序列。\n"
     "• test_agent_tools_store.py：测试智能体工具箱（Tools）与状态存储器（Store）的数据读写与工具执行。\n"
     "• test_agent_api.py：测试智能体对外暴露的后端 Web API 接口（/api/agent/*）。"),
    ("二、 混合检索与 RAG 引擎测试（Retrieval & RAG 专项）",
     "测试双引擎检索算法、排名融合与问答管道：\n"
     "• test_hybrid_retriever.py：测试 Hybrid 混合检索器，验证 TF-IDF 字面分与 BGE 语义向量分的加权融合算法（RRF/加权平均）。\n"
     "• test_rag_pipeline.py：测试完整的 RAG 问答主流水线（端到端检索 → 证据筛选 → 提示词组装 → 生成）。\n"
     "• test_composite_pipeline.py：测试复合问题问答流（例如同时询问两项不同业务的交叉问题）。\n"
     "• test_canteen_location.py：专项测试食堂、宿舍等地理位置型问答的召回与定位精度。\n"
     "• test_topic_coverage.py：测试校本规章各个主题（奖助学金、选课、毕业、图书等）的检索覆盖率。"),
    ("三、 意图理解、槽位抽取与证据链测试（NLU & Evidence 专项）",
     "测试对用户错别字、倒装口语的容错以及证据片段的真实性校验：\n"
     "• test_query_understanding.py：测试查询理解模块（分词、实体识别、意图分类）。\n"
     "• test_fact_slots.py：测试事实槽位提取器（从非结构化文本中精准抠出电话、时间、地点、金额等硬事实字段）。\n"
     "• test_evidence.py：测试证据重排序与证据段落提取，确保回答只引用高置信度官方段落。\n"
     "• test_article_reading.py：测试长篇规章制度文档（如章程、条例）的细粒度段落切分与局部阅读理解。"),
    ("四、 接口契约与大模型客户端测试（API & Client 专项）",
     "测试前后端数据交互协议和大模型接口适配：\n"
     "• test_chat_schemas.py：基于 Pydantic 的接口契约测试（测试请求体字段校验、非法参数拦截、格式异常报错）。\n"
     "• test_deepseek_client.py：测试大模型 API 客户端封装（超时重试机制、Token 统计、流式输出解析与断网模拟）。\n"
     "• test_custom_ai_and_history.py：测试用户自定义大模型端点（支持切换第三方 API 或本地 Ollama）与历史记录管理。"),
    ("五、 底层数据加载与向量缓存测试（Data & Cache 专项）",
     "测试离线数据处理与存储性能：\n"
     "• test_document_loader.py：测试校本 Markdown/文本知识库的清洗、切分与元数据加载。\n"
     "• test_embedding_cache.py：测试 BGE 向量离线缓存（.npy 和哈希校验），确保服务器启动时秒级加载已有向量，不重复编码。\n"
     "• test_db_manager.py：测试 SQLite 本地数据库的管理（问答日志存储、反馈打标、历史会话持久化）。\n"
     "• test_kb_display.py：测试前端展示的知识库文档数与语义块统计是否与底层实际加载量一致。"),
    ("六、 评测与断言逻辑测试（Evaluation 专项）",
     "• test_answer_evaluate.py：测试评测脚本自身的准确性判定算法（如对比模型答案与标准答案时的 ROUGE/BLEU 评分及事实一致性打分逻辑）。"),
    ("七、 前端 UI 自动化测试（Node.js 端）",
     "• test_chat_ui.cjs：使用 Node.js 针对对话界面的前端交互、气泡渲染、历史记录加载进行的自动化测试。\n"
     "• test_agent_ui.cjs：针对智能体多步执行轨迹折叠面板、工具调用卡片等 UI 组件进行的渲染测试。")
]

for title, content in sections_data:
    p_title = ref_p.insert_paragraph_before()
    r = p_title.add_run(title)
    r.bold = True
    
    p_content = ref_p.insert_paragraph_before(content)

doc.save(doc_path)
print("Docx successfully updated with exact heading: tests/ 目录下 25 份测试脚本说明：")
